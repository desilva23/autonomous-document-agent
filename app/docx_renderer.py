"""Render the agent's structured output into a styled Word document (python-docx).

Layout: cover page -> document control + contents -> numbered sections (executive
summary as a callout, data tables with a branded header row) -> assumptions,
conflicts and open questions -> header/footer with page X of Y.
"""
import itertools
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

from .schemas import PlannedTask, RequestAnalysis, ReviewRound, SectionContent, TableData

NAVY, ACCENT, INK, MUTED = "1F3A5F", "2E75B6", "262626", "6B7280"
ZEBRA, CALLOUT, RULE, LABEL = "F3F6FA", "EAF2FB", "C9D3E0", "E8EDF4"
FONT = "Calibri"
MARGIN = Cm(2.2)
TEXT_WIDTH = Emu(Cm(21.0) - 2 * MARGIN)

_BOLD = re.compile(r"\*\*(.+?)\*\*")
_NUMERIC = re.compile(r"^[~≈]?[$€£₹]?\s?-?[\d,.]+\s?(%|[kKmMbB]|hrs?|days?|weeks?)?$")
# Elements that must follow pBdr in pPr / spacing in rPr (Word rejects out-of-order XML).
_PPR_AFTER_BDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                  "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
                  "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                  "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
                  "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
_RPR_AFTER_SPACING = ("w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr",
                      "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout",
                      "w:specVanish", "w:oMath")
_TCPR_ORDER = ["tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
               "textDirection", "tcFitText", "vAlign", "hideMark"]


def render_document(analysis: RequestAnalysis, tasks: list[PlannedTask], sections: dict[str, SectionContent],
                    rounds: list[ReviewRound], path: Path, verification: str = "") -> Path:
    doc = Document()
    _page_setup(doc)
    _styles(doc)

    ordered = [(t, sections[t.id]) for t in tasks if t.id in sections]
    has_appendix = bool(analysis.assumptions or analysis.conflicts or analysis.missing_information)
    headings = [t.title for t, _ in ordered] + (["Assumptions, Decisions & Open Questions"] if has_appendix else [])

    _cover(doc, analysis)
    _control_and_contents(doc, analysis, headings, rounds, verification)

    tables = itertools.count(1)
    for n, (task, content) in enumerate(ordered, 1):
        _section(doc, n, task, content, tables)
    if has_appendix:
        _assumptions(doc, len(ordered) + 1, analysis, tables)

    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    end.paragraph_format.space_before = Pt(24)
    _run(end, "— End of document —", size=9, color=MUTED, italic=True)

    _header_footer(doc, analysis)
    props = doc.core_properties
    props.title, props.subject, props.author = analysis.title, analysis.document_type, "Autonomous Document Agent"

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


# ---------------------------------------------------------------- page & styles

def _page_setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4
    sec.left_margin = sec.right_margin = MARGIN
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    sec.different_first_page_header_footer = True  # clean cover page


def _set_font(style_element, name=FONT):
    fonts = style_element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), name)


def _styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    _set_font(normal.element)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, before, after in (("Heading 1", 15, NAVY, 20, 6), ("Heading 2", 12, ACCENT, 12, 4)):
        st = styles[name]
        _set_font(st.element)
        st.font.size, st.font.bold = Pt(size), True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(before), Pt(after)
        st.paragraph_format.keep_with_next = True

    bullet = styles["List Bullet"]
    bullet.paragraph_format.space_after = Pt(3)
    bullet.paragraph_format.line_spacing = 1.1


# ---------------------------------------------------------------- text helpers

def _run(p, text, size=None, color=None, bold=False, italic=False):
    r = p.add_run(text)
    r.bold, r.italic = bold or None, italic or None
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def _rich(p, text, size=None, color=None):
    """Add text to a paragraph, honouring **bold** spans."""
    text, pos = text.strip(), 0
    for m in _BOLD.finditer(text):
        if m.start() > pos:
            _run(p, text[pos:m.start()], size, color)
        _run(p, m.group(1), size, color, bold=True)
        pos = m.end()
    if pos < len(text):
        _run(p, text[pos:].replace("**", ""), size, color)
    return p


def _bullet(container, text, size=None):
    p = container.add_paragraph(style="List Bullet")
    m = re.match(r"^([^:*]{2,48}):\s+(.+)$", text)
    if m and len(m.group(1).split()) <= 6:  # "Label: detail" -> bold label
        _run(p, m.group(1) + ": ", size, bold=True)
        _rich(p, m.group(2), size)
    else:
        _rich(p, text, size)
    return p


def _bottom_rule(p, color=ACCENT, size=12, space=4):
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", str(size)), ("space", str(space)), ("color", color)):
        bottom.set(qn(f"w:{k}"), v)
    bdr.append(bottom)
    ppr.insert_element_before(bdr, *_PPR_AFTER_BDR)


def _letter_spacing(run, twips):
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(twips))
    run._r.get_or_add_rPr().insert_element_before(spacing, *_RPR_AFTER_SPACING)


def _field(p, instr, size=9, color=MUTED):
    """Insert a Word field (PAGE, NUMPAGES) that Word evaluates on open."""
    def fld_char(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        _run(p, "", size, color)._r.append(el)

    fld_char("begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = f" {instr} "
    _run(p, "", size, color)._r.append(code)
    fld_char("separate")
    _run(p, "1", size, color)
    fld_char("end")


# ---------------------------------------------------------------- table helpers

def _tc_set(cell, element):
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(element.tag):
        tcpr.remove(old)
    tcpr.append(element)
    children = sorted(tcpr, key=lambda e: _TCPR_ORDER.index(e.tag.split("}")[1])
                      if e.tag.split("}")[1] in _TCPR_ORDER else 99)
    for child in children:  # Word is strict about child order inside tcPr
        tcpr.remove(child)
        tcpr.append(child)


def _set_widths(table, widths):
    """Pin table, grid and cell widths so every renderer (Word, Pages, Google Docs) agrees."""
    table.autofit = False
    tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(sum(widths) / 635)))  # EMU -> twips
    for col, w in zip(table.columns, widths):
        col.width = w
        for cell in col.cells:
            cell.width = w


def _shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    _tc_set(cell, shd)


def _borders(cell, color=RULE, size=4, **overrides):
    """All four edges; pass e.g. left=(ACCENT, 24) or top=None to customise/remove."""
    el = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        spec = overrides.get(edge, (color, size))
        b = OxmlElement(f"w:{edge}")
        if spec is None:
            b.set(qn("w:val"), "nil")
        else:
            for k, v in (("val", "single"), ("sz", str(spec[1])), ("space", "0"), ("color", spec[0])):
                b.set(qn(f"w:{k}"), v)
        el.append(b)
    _tc_set(cell, el)


def _margins(cell, top=80, side=120):
    el = OxmlElement("w:tcMar")
    for edge, w in (("top", top), ("left", side), ("bottom", top), ("right", side)):
        m = OxmlElement(f"w:{edge}")
        m.set(qn("w:w"), str(w))
        m.set(qn("w:type"), "dxa")
        el.append(m)
    _tc_set(cell, el)


def _cell_text(cell, text, size=9.5, color=None, bold=False, align=None):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    if align is not None:
        p.alignment = align
    if bold:
        _run(p, text.replace("**", ""), size, color, bold=True)
    else:
        _rich(p, text, size, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _callout(doc, fill=CALLOUT, bar=ACCENT):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_widths(t, [TEXT_WIDTH])
    cell = t.cell(0, 0)
    _shade(cell, fill)
    _borders(cell, top=None, right=None, bottom=None, left=(bar, 24))
    _margins(cell, top=140, side=220)
    return cell


def _data_table(doc, table: TableData, number: int):
    if table.caption:
        cap = doc.add_paragraph()
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.space_before, cap.paragraph_format.space_after = Pt(6), Pt(4)
        _run(cap, f"Table {number}. ", 9, NAVY, bold=True)
        _run(cap, table.caption, 9, MUTED, italic=True)

    cols = table.columns
    t = doc.add_table(rows=1, cols=len(cols))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Column widths proportional to content length, so description columns get the room.
    lengths = [max([len(c)] + [len(r[i]) for r in table.rows]) for i, c in enumerate(cols)]
    weights = [min(max(n, 6), 45) for n in lengths]
    widths = [int(TEXT_WIDTH * w / sum(weights)) for w in weights]
    # No column narrower than its longest word (so "Competitor" never breaks mid-word); the space comes
    # proportionally from the columns that have room to spare.
    longest_word = [max(len(w) for cell in [c] + [r[i] for r in table.rows] for w in (cell.split() or [""]))
                    for i, c in enumerate(cols)]
    mins = [int(Cm(0.2) * n + Cm(0.45)) for n in longest_word]
    deficit = sum(max(0, m - w) for m, w in zip(mins, widths))
    if deficit:
        spare = {i: w - m for i, (w, m) in enumerate(zip(widths, mins)) if w > m}
        widths = [max(w, m) for w, m in zip(widths, mins)]
        for i, room in spare.items():
            widths[i] -= int(min(deficit, sum(spare.values())) * room / sum(spare.values()))
        widths = [int(w * TEXT_WIDTH / sum(widths)) for w in widths]  # still fits the page

    header = t.rows[0]
    trpr = header._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")  # repeat header row on page breaks
    repeat.set(qn("w:val"), "true")
    trpr.append(repeat)
    for i, name in enumerate(cols):
        cell = header.cells[i]
        _cell_text(cell, name, 9.5, "FFFFFF", bold=True)
        _shade(cell, NAVY)
        _borders(cell, color=NAVY)
        _margins(cell)

    for ri, row in enumerate(table.rows):
        cells = t.add_row().cells
        total = row[0].strip().lower().startswith(("total", "grand total", "subtotal"))
        for i, value in enumerate(row):
            numeric = bool(_NUMERIC.match(value.strip())) and i > 0
            _cell_text(cells[i], value, 9.5, bold=total, align=WD_ALIGN_PARAGRAPH.RIGHT if numeric else None)
            if ri % 2:
                _shade(cells[i], ZEBRA)
            _borders(cells[i])
            _margins(cells[i])

    _set_widths(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ---------------------------------------------------------------- page blocks

def _cover(doc, a: RequestAnalysis):
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(150)
    _letter_spacing(_run(kicker, a.document_type.upper(), 11, ACCENT, bold=True), 40)

    title = doc.add_paragraph()
    title.paragraph_format.line_spacing = 1.0
    title.paragraph_format.space_after = Pt(8)
    _run(title, a.title, 28, NAVY, bold=True)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(28)
    _run(rule, a.subtitle or a.objective, 13, MUTED)
    _bottom_rule(rule, ACCENT, 18, 12)

    meta = doc.add_table(rows=0, cols=2)
    for label, value in (("Prepared for", a.audience), ("Prepared by", "Autonomous Document Agent"),
                         ("Date", date.today().strftime("%d %B %Y")), ("Version", "1.0 — Draft for review")):
        cells = meta.add_row().cells
        _cell_text(cells[0], label.upper(), 8.5, MUTED, bold=True)
        _cell_text(cells[1], value, 10.5, INK)
        for c in cells:
            _margins(c, top=60, side=60)
    _set_widths(meta, [Cm(3.6), TEXT_WIDTH - Cm(3.6)])

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _control_and_contents(doc, a: RequestAnalysis, headings: list[str], rounds: list[ReviewRound],
                          verification: str = ""):
    doc.add_heading("Document Control", level=2)
    quality = "Not reviewed"
    if rounds:
        revised = sum(len(r.revised_sections) + len(r.added_sections) for r in rounds)
        shipped = next((r for r in rounds if r.shipped), rounds[-1])
        quality = (f"{shipped.score}/10 after automated self-review "
                   f"(scores {' → '.join(f'{r.score}' for r in rounds)}; {revised} section revision{'s' * (revised != 1)})")
    rows = (("Document", a.title), ("Type", a.document_type), ("Audience", a.audience),
            ("Objective", a.objective or "—"), ("Status", "Draft v1.0 — for stakeholder review"),
            ("Quality gate", quality)) + ((("Verification", verification),) if verification else ())
    t = doc.add_table(rows=0, cols=2)
    for label, value in rows:
        cells = t.add_row().cells
        _cell_text(cells[0], label, 9.5, NAVY, bold=True)
        _cell_text(cells[1], value, 9.5)
        _shade(cells[0], LABEL)
        for c in cells:
            _borders(c)
            _margins(c)
    _set_widths(t, [Cm(3.6), TEXT_WIDTH - Cm(3.6)])

    doc.add_heading("Contents", level=2).paragraph_format.space_before = Pt(22)
    for i, h in enumerate(headings, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
        _run(p, f"{i}\t", 10.5, ACCENT, bold=True)
        _run(p, h, 10.5)
        _bottom_rule(p, RULE, 4, 3)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _section(doc, n: int, task: PlannedTask, c: SectionContent, tables):
    doc.add_heading(f"{n}. {task.title}", level=1)

    if task.action == "executive_summary":
        cell = _callout(doc)
        first = True
        for para in c.paragraphs:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            p.paragraph_format.space_after = Pt(6)
            _rich(p, para)
        if c.bullets:
            label = cell.add_paragraph()
            label.paragraph_format.space_before, label.paragraph_format.space_after = Pt(4), Pt(2)
            _run(label, "KEY POINTS", 8.5, ACCENT, bold=True)
            for b in c.bullets:
                _bullet(cell, b)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return

    for para in c.paragraphs:
        _rich(doc.add_paragraph(), para)
    if c.table:
        _data_table(doc, c.table, next(tables))
    for b in c.bullets:
        _bullet(doc, b)


def _assumptions(doc, n: int, a: RequestAnalysis, tables):
    doc.add_heading(f"{n}. Assumptions, Decisions & Open Questions", level=1)
    _rich(doc.add_paragraph(),
          "The request did not specify every detail. To deliver a complete draft without blocking on clarification, "
          "the agent made the decisions below. Each should be confirmed with the relevant stakeholder before this "
          "document is finalised.")
    if a.conflicts:
        doc.add_heading("Conflicting requirements and how they were resolved", level=2)
        _data_table(doc, TableData(caption="Conflict resolution log", columns=["Conflict", "Resolution adopted"],
                                   rows=[[c.issue, c.resolution or "Flagged for stakeholder decision"]
                                         for c in a.conflicts]), next(tables))
    if a.assumptions:
        doc.add_heading("Assumptions made", level=2)
        for x in a.assumptions:
            _bullet(doc, x)
    if a.missing_information:
        doc.add_heading("Open questions for stakeholders", level=2)
        for x in a.missing_information:
            _bullet(doc, x)


def _header_footer(doc, a: RequestAnalysis):
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(hp, a.title, 8.5, MUTED)
    _bottom_rule(hp, RULE, 4, 4)

    fp = sec.footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    _run(fp, f"{a.document_type}  ·  Confidential — draft for review", 8.5, MUTED)
    _run(fp, "\tPage ", 8.5, MUTED)
    _field(fp, "PAGE", 8.5)
    _run(fp, " of ", 8.5, MUTED)
    _field(fp, "NUMPAGES", 8.5)
